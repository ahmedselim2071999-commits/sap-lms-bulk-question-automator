from flask import Flask, request, render_template, send_file, jsonify
from pathlib import Path
from werkzeug.utils import secure_filename
import tempfile, os, re, copy
from openpyxl import load_workbook

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024

SUPPORTED_TYPES = {
    "MULTI_CHOICE_SINGLE_ANSWER",
    "MULTI_CHOICE_MULTIPLE_ANSWER",
    "TRUE_FALSE",
    "ORDERING",
    "FILL_IN_THE_BLANK",
}
ASSESSMENT_CODES = {
    "Pre Assessment": "PA",
    "Knowledge Check": "KC",
    "Post Assessment": "POST",
}
LETTER_RE = re.compile(r"^([A-Za-z])[\.\)\:\-]\s*(.+)$")
ANSWER_RE = re.compile(r"^(?:ANSWER|CORRECT ANSWER|ANS)\s*:\s*(.+)$", re.I)
TYPE_RE = re.compile(r"^\s*(?:TYPE|QUESTION TYPE)\s*:\s*(.+)$", re.I)

def norm_type(v):
    v = v.strip().upper().replace(" ", "_").replace("-", "_")
    aliases = {
        "SINGLE": "MULTI_CHOICE_SINGLE_ANSWER",
        "SINGLE_ANSWER": "MULTI_CHOICE_SINGLE_ANSWER",
        "MULTIPLE": "MULTI_CHOICE_MULTIPLE_ANSWER",
        "MULTIPLE_ANSWER": "MULTI_CHOICE_MULTIPLE_ANSWER",
        "TRUEFALSE": "TRUE_FALSE",
        "ORDER": "ORDERING",
        "FILL_BLANK": "FILL_IN_THE_BLANK",
        "FILL_IN_BLANK": "FILL_IN_THE_BLANK",
    }
    v = aliases.get(v, v)
    if v not in SUPPORTED_TYPES:
        raise ValueError(f"Unsupported question type: {v}")
    return v

def read_text(upload):
    name = upload.filename.lower()
    if name.endswith(".docx"):
        from docx import Document
        data = upload.read()
        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
            f.write(data); p = f.name
        try:
            doc = Document(p)
            lines = [x.text.rstrip() for x in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            lines.extend(cell.text.splitlines())
                    lines.append("")
            return lines
        finally:
            os.unlink(p)
    return upload.read().decode("utf-8-sig").splitlines()

def parse_questions(lines, default_type):
    qs, cur, answer, qtype = [], None, None, None
    def flush():
        nonlocal cur, answer, qtype
        if cur is None: return
        text = cur["text"].strip()
        choices = cur["choices"]
        typ = norm_type(qtype or default_type)
        if typ == "FILL_IN_THE_BLANK":
            if not answer: raise ValueError(f"Missing ANSWER for: {text}")
            choices = [{"text": answer.strip()}]
            correct = {0}
        else:
            if not choices: raise ValueError(f"No choices found for: {text}")
            if not answer: raise ValueError(f"Missing ANSWER for: {text}")
            tokens = [x.strip() for x in re.split(r"[,;|]", answer) if x.strip()]
            correct = set()
            for tok in tokens:
                if re.fullmatch(r"[A-Za-z]", tok):
                    idx = ord(tok.upper()) - 65
                    if idx >= len(choices): raise ValueError(f"Answer {tok} is outside choices: {text}")
                    correct.add(idx)
                else:
                    hits = [i for i,c in enumerate(choices) if c["text"].strip().casefold()==tok.casefold()]
                    if not hits: raise ValueError(f"Could not match answer '{tok}': {text}")
                    correct.add(hits[0])
            if typ == "MULTI_CHOICE_SINGLE_ANSWER" and len(correct) != 1:
                raise ValueError(f"Single-answer question needs exactly one correct answer: {text}")
            if typ == "MULTI_CHOICE_MULTIPLE_ANSWER" and not correct:
                raise ValueError(f"Multiple-answer question needs at least one correct answer: {text}")
        qs.append({"text":text,"choices":choices,"correct":correct,"type":typ})
        cur=answer=qtype=None
    for raw in lines:
        line=raw.strip()
        if not line:
            if cur is not None and answer is not None: flush()
            continue
        m=ANSWER_RE.match(line)
        if m:
            if cur is None: raise ValueError("ANSWER appeared before a question.")
            answer=m.group(1).strip(); continue
        m=TYPE_RE.match(line)
        if m:
            qtype=m.group(1).strip(); continue
        m=LETTER_RE.match(line)
        if m:
            if cur is None: raise ValueError("Choice appeared before a question.")
            cur["choices"].append({"text":m.group(2).strip()}); continue
        if cur is not None and answer is not None: flush()
        if cur is None: cur={"text":line,"choices":[]}
        else: cur["text"] += " " + line
    flush()
    return qs

def create_xlsx(template_bytes, questions, course, stage, locale, domain, start=1):
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        f.write(template_bytes); template=f.name
    try:
        wb=load_workbook(template)
        if "Question" not in wb.sheetnames:
            raise ValueError("The template does not contain a 'Question' sheet.")
        ws=wb["Question"]
        headers={ws.cell(1,c).value:c for c in range(1,ws.max_column+1)}
        required=["Question ID (*required)","Locale ID (*required)","Domain ID (*required)",
                  "Variant Number (*required)","Question Type (*required)","Question Text (*required)",
                  "Answer Choice Number (*required)","Answer Choice Value","Is Correct (*required)"]
        miss=[x for x in required if x not in headers]
        if miss: raise ValueError("Missing template columns: "+", ".join(miss))
        style_row=2
        if ws.max_row>1: ws.delete_rows(2,ws.max_row-1)
        row=2
        prefix=re.sub(r"[^A-Za-z0-9_.-]+","_",course.strip()).strip("_")
        code=ASSESSMENT_CODES.get(stage, stage.upper().replace(" ","_"))
        for qi,q in enumerate(questions,start=start):
            qid=f"{prefix}_{code}_{qi:03d}"
            choices=q["choices"]
            if q["type"]=="TRUE_FALSE" and not choices:
                choices=[{"text":"True"},{"text":"False"}]
            for ci,ch in enumerate(choices,start=1):
                if row>2:
                    for c in range(1,ws.max_column+1):
                        src=ws.cell(style_row,c); dst=ws.cell(row,c)
                        if src.has_style: dst._style=copy.copy(src._style)
                        dst.alignment=copy.copy(src.alignment)
                        dst.protection=copy.copy(src.protection)
                        dst.number_format=src.number_format
                vals={
                    "Question ID (*required)":qid,"Locale ID (*required)":locale,
                    "Domain ID (*required)":domain,"Variant Number (*required)":1,
                    "Question Type (*required)":q["type"],"Question Text (*required)":q["text"],
                    "Answer Choice Number (*required)":ci,"Answer Choice Value":ch["text"],
                    "Is Correct (*required)":"Y" if ci-1 in q["correct"] else "N"
                }
                if "Point Value" in headers: vals["Point Value"]=1
                if "Available for Exams" in headers: vals["Available for Exams"]="Y"
                for k,v in vals.items(): ws.cell(row,headers[k]).value=v
                row+=1
        for sname in ("Feedback","Objectives"):
            if sname in wb.sheetnames and wb[sname].max_row>1:
                wb[sname].delete_rows(2,wb[sname].max_row-1)
        out=tempfile.NamedTemporaryFile(suffix=".xlsx",delete=False)
        out.close(); wb.save(out.name)
        return out.name
    finally:
        os.unlink(template)

@app.route("/")
def index():
    return render_template("index.html")

@app.post("/generate")
def generate():
    try:
        source=request.files.get("source")
        template=request.files.get("template")
        if not source or not template:
            return jsonify(ok=False,error="Upload both the question file and SAP template."),400
        course=request.form.get("course","").strip()
        stage=request.form.get("stage","Knowledge Check")
        locale=request.form.get("locale","en_US").strip() or "en_US"
        domain=request.form.get("domain","PUBLIC").strip() or "PUBLIC"
        default_type=request.form.get("qtype","MULTI_CHOICE_SINGLE_ANSWER")
        if not course: return jsonify(ok=False,error="Course code is required."),400
        qs=parse_questions(read_text(source),default_type)
        if not qs: return jsonify(ok=False,error="No questions detected."),400
        result=create_xlsx(template.read(),qs,course,stage,locale,domain)
        return send_file(result,as_attachment=True,download_name=f"{course}_{ASSESSMENT_CODES.get(stage,'ASSESSMENT')}_SAP_Bulk.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except Exception as e:
        return jsonify(ok=False,error=str(e)),400

if __name__=="__main__":
    app.run(host="127.0.0.1",port=5000,debug=False)
